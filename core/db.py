# -*- coding: utf-8 -*-
"""SQLite 客户库：连接、建表、去重键、upsert、统计（原 `客户开发工具.py` 第 375–530 行整段搬入）。

依赖：`core.paths`（路径与去重正则）+ `core.state`（log）。
`log` 之所以必须来自 `core.state` 而不是入口：`init_customer_db()` 要调它，
若 log 留在入口，这里就得反过来 import 入口 → 循环导入。
"""
import os
import re
import sqlite3
import time
from contextlib import closing
from urllib.parse import urlparse

from openpyxl import load_workbook

from core import paths
from core.paths import DB_DIR, DB_FILE, MAPS_PLACE_ID_RE, _app_dir
from core.state import log


# ================= SQLite 客户库 =================
def db_path():
    base_dir = _app_dir()          # 程序根（不是「本文件所在目录」，见 core/paths.py 的 _app_dir 注释）
    db_dir = os.path.join(base_dir, DB_DIR)
    os.makedirs(db_dir, exist_ok=True)   # DB 文件夹不存在时自动创建
    return os.path.join(db_dir, DB_FILE)


def maps_place_key(url):
    """从谷歌地图 place 链接里提取稳定去重键（结尾的 0x...:0x... 十六进制段）。
    比商家名可靠（同名商家多），也比 place_id 好取（place_id 还得额外解析一次）。
    取不到返回空串。"""
    m = MAPS_PLACE_ID_RE.search(url or '')
    return ('maps:' + m.group(1).lower()) if m else ''


def normalize_domain(url):
    """把网址规整为稳定的公司去重键。"""
    try:
        raw = (url or '').strip().lower()
        if '@' in raw and '://' not in raw:
            return ''
        # 谷歌地图 place 链接：没官网的商家会落到这里。必须给个稳定键，
        # 否则 netloc 全是 google.com —— 既被搜索引擎黑名单拦掉，又互相覆盖成一条。
        mk = maps_place_key(raw)
        if mk:
            return mk
        netloc = urlparse(raw if '://' in raw else 'https://' + raw).netloc
        netloc = netloc.split(':', 1)[0].removeprefix('www.')
        return netloc
    except Exception:
        return ''


def normalize_email(email):
    return (email or '').strip().lower()


def db_connect():
    conn = sqlite3.connect(db_path(), timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA journal_mode=WAL')
    return conn


def init_customer_db():
    """初始化客户库，并把已有评分 Excel 导入一次。"""
    with closing(db_connect()) as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS customers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            company_name TEXT NOT NULL DEFAULT '',
            website TEXT NOT NULL DEFAULT '',
            website_key TEXT UNIQUE,
            email TEXT DEFAULT '',
            phone TEXT DEFAULT '',
            contact_url TEXT DEFAULT '',
            country TEXT DEFAULT '',
            industry TEXT DEFAULT '',
            score INTEGER,
            rating TEXT DEFAULT '',
            snippet TEXT DEFAULT '',
            search_formula TEXT DEFAULT '',
            source TEXT DEFAULT '谷歌搜索',
            status TEXT DEFAULT '未联系',
            next_follow_up TEXT DEFAULT '',
            notes TEXT DEFAULT '',
            first_seen_at TEXT NOT NULL DEFAULT (datetime('now','localtime')),
            last_seen_at TEXT NOT NULL DEFAULT (datetime('now','localtime')),
            verified_at TEXT DEFAULT ''
        )''')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_customers_email ON customers(email)')
        # 老库升级：补齐后加的列（已有 DB 不会因缺列报错）
        cols = [r[1] for r in conn.execute('PRAGMA table_info(customers)')]
        if 'industry' not in cols:
            conn.execute("ALTER TABLE customers ADD COLUMN industry TEXT DEFAULT ''")
        conn.commit()
    # 只在库为空时迁移旧结果，保证重复运行不会重复导入。
    with closing(db_connect()) as conn:
        empty = conn.execute('SELECT 1 FROM customers LIMIT 1').fetchone() is None
    if empty and os.path.exists(paths.ENRICH_FILE):     # ⚠️ 须按属性读，才是最新值
        try:
            wb = load_workbook(paths.ENRICH_FILE, read_only=True)
            ws = wb.active
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r and len(r) > 1 and r[1]:
                    db_upsert_customer({
                        'company_name': r[0], 'website': r[1],
                        'email': r[2] or '', 'phone': r[3] or '',
                        'contact_url': r[4] or '', 'score': r[5],
                        'rating': r[6] or '', 'snippet': r[7] or '',
                        'search_formula': r[8] or '',
                        'country': r[9] if len(r) > 9 else '',
                    })
            wb.close()
            log('✅ 已将现有评分 Excel 导入 SQLite 客户库')
        except Exception as ex:
            log(f'⚠️ 旧 Excel 导入客户库失败：{type(ex).__name__}: {ex}')


def db_upsert_customer(data):
    """新增或合并客户；空值不会覆盖数据库里已有的有效值。"""
    website = (data.get('website') or '').strip()
    website_key = normalize_domain(website)
    email = normalize_email(data.get('email', ''))
    if not website_key and not email:
        return False
    fields = {
        'company_name': data.get('company_name', ''), 'website': website,
        'email': email, 'phone': data.get('phone', ''),
        'contact_url': data.get('contact_url', ''), 'country': data.get('country', ''),
        'industry': data.get('industry', ''),
        'score': data.get('score'), 'rating': data.get('rating', ''),
        'snippet': data.get('snippet', ''), 'search_formula': data.get('search_formula', ''),
        'source': data.get('source', '谷歌搜索'), 'notes': data.get('notes', ''),
    }
    with closing(db_connect()) as conn:
        row = None
        if website_key:
            row = conn.execute('SELECT * FROM customers WHERE website_key=?', (website_key,)).fetchone()
        if row is None and email:
            row = conn.execute('SELECT * FROM customers WHERE lower(email)=?', (email,)).fetchone()
        if row is None:
            conn.execute('''INSERT INTO customers
                (company_name,website,website_key,email,phone,contact_url,country,industry,score,rating,
                 snippet,search_formula,source,notes)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (fields['company_name'], fields['website'], website_key, fields['email'],
                 fields['phone'], fields['contact_url'], fields['country'], fields['industry'],
                 fields['score'], fields['rating'], fields['snippet'], fields['search_formula'],
                 fields['source'], fields['notes']))
            conn.commit()
            return True
        updates = []
        values = []
        for key, value in fields.items():
            if value not in ('', None) and (key != 'score' or isinstance(value, (int, float))):
                updates.append(f'{key}=?')
                values.append(value)
        # ---- 键自愈：让 website_key 恒等于 normalize_domain(website) ----
        # 命中的这条可能是**按 email 兜底**匹配上的（两家站共用同一个邮箱时会被合并），
        # 此时它的 website_key 还是旧站域名 —— 键一旦和网址错位，
        # db_has_customer('新站') 就永远查不到 → 同一个站每次采集都被判「新客户」、
        # 反复写进结果 Excel，而库里始终只有一条（实测：megazone.mx 在一份表里出现 3 次）。
        # ⚠️ 改键前先确认新键没被别的行占用，否则这条 UPDATE 会撞 UNIQUE 约束。
        #    按上面的查找顺序（先按键、再按邮箱），此处的 website_key 一定没被占用；
        #    这个判断只是防将来把两步调换顺序时静默炸掉。
        if website_key and website_key != (row['website_key'] or ''):
            taken = conn.execute('SELECT id FROM customers WHERE website_key=?',
                                 (website_key,)).fetchone()
            if taken is None:
                if row['website_key']:
                    log(f'🔧 键自愈：id={row["id"]} 的 website_key '
                        f'{row["website_key"]!r} → {website_key!r}（与网址对齐）')
                updates.append('website_key=?')
                values.append(website_key)
            else:
                log(f'⚠️ id={row["id"]} 的网址算出键 {website_key!r}，但已被 id={taken["id"]} 占用，'
                    f'键保持 {row["website_key"]!r} 不变')
        updates.append("last_seen_at=datetime('now','localtime')")
        values.append(row['id'])
        conn.execute(f"UPDATE customers SET {', '.join(updates)} WHERE id=?", values)
        conn.commit()
        return False


def db_stats():
    init_customer_db()
    with closing(db_connect()) as conn:
        total = conn.execute('SELECT COUNT(*) FROM customers').fetchone()[0]
        verified = conn.execute("SELECT COUNT(*) FROM customers WHERE score IS NOT NULL").fetchone()[0]
    return {'total': total, 'verified': verified, 'file': DB_DIR + '/' + DB_FILE}


# ================= 去重键一致性（website_key == normalize_domain(website)）=================
# 这是全库最重要的一条不变量：db_has_customer()（判断「要不要跳过」）**只**按 website_key 查，
# 而 db_upsert_customer()（写入）是「按 website_key 查不到、再按 email 兜底合并」——
# 两者口径本来就不完全一致。一旦某行的键和网址错位，那个网址就会被永久判成「新客户」，
# 每采一次就往结果 Excel 里多写一行，而库里始终只有一条。
def db_fix_website_keys(apply=False):
    """把所有「website_key != normalize_domain(website)」的行修回一致。

    返回 (fixed, conflict) 两个列表，元素是 (id, 旧键, 应改成的键, ...)：
      · fixed    —— 能安全修正的（apply=True 时真的写库）
      · conflict —— 目标键已被**别的行**占用的，不硬改（会撞 UNIQUE），留给人判断

    apply=False 是纯只读 dry-run（回归脚本用这个模式断言「真实库里 conflict/fixed 都为空」）。
    """
    init_customer_db()
    fixed, conflict = [], []
    with closing(db_connect()) as conn:
        rows = list(conn.execute('SELECT id, website_key, website FROM customers'))
        holder = {(r['website_key'] or ''): r['id'] for r in rows}
        for r in rows:
            want = normalize_domain(r['website'] or '')
            if not want or want == (r['website_key'] or ''):
                continue
            other = holder.get(want)
            if other is not None and other != r['id']:
                conflict.append((r['id'], r['website_key'], want, other))
                continue
            fixed.append((r['id'], r['website_key'], want, r['website']))
            if apply:
                conn.execute('UPDATE customers SET website_key=? WHERE id=?', (want, r['id']))
        if apply:
            conn.commit()
    if apply:
        for i, old, new, _w in fixed:
            log(f'🔧 键修正：id={i} {old!r} → {new!r}')
        for i, old, new, other in conflict:
            log(f'⚠️ 键冲突未改：id={i} {old!r} → {new!r} 已被 id={other} 占用')
    return fixed, conflict


# ================= 删除客户（含删除前自动备份）=================
# 「本轮采集明细」每行一个 ✕：把不想要的条目从本轮清单里拿掉，同时把库里那条也删掉。
# 这是**不可逆**操作，所以删之前先做一份 WAL 安全的一致性备份。
_DELETE_BACKUP_MIN_INTERVAL = 300   # 秒。连点删多行时只留一份，免得把 DB/backups 撑爆
_last_delete_backup_at = 0.0


def _backups_dir():
    """备份目录 = 库文件同级的 backups/。
    刻意由 db_path() 推导：回归脚本只要打桩了 core.db.db_path，备份就自动落进临时目录，
    护栏不必再为备份目录单独补一处打桩（漏打就会往真实 DB/backups/ 里塞测试垃圾）。"""
    return os.path.join(os.path.dirname(db_path()), 'backups')


def backup_db_before_delete():
    """删除客户前做一次一致性备份，返回备份文件绝对路径；被节流或失败返回 ''。
    ⚠️ 必须用 sqlite3 的在线 backup()，**禁止 shutil.copy2** —— 库是 WAL 模式，裸拷贝会漏掉
       -wal 里尚未 checkpoint 的写入，备出来的可能缺最新数据（同 重标行业.py 的做法）。
    ⚠️ 备份失败**不阻断**删除（用户按的按钮是「删除」，不是「备份」），只记日志、返回 ''。"""
    global _last_delete_backup_at
    now = time.time()
    if now - _last_delete_backup_at < _DELETE_BACKUP_MIN_INTERVAL:
        return ''
    try:
        d = _backups_dir()
        os.makedirs(d, exist_ok=True)
        dst_path = os.path.join(d, '客户库_删除前_' + time.strftime('%Y%m%d_%H%M%S') + '.sqlite3')
        src = sqlite3.connect(db_path())
        dst = sqlite3.connect(dst_path)
        try:
            with dst:
                src.backup(dst)
        finally:
            dst.close()
            src.close()
        _last_delete_backup_at = now
        log('🗄️ 删除前已备份客户库：' + dst_path)
        return dst_path
    except Exception as ex:
        log(f'⚠️ 删除前备份客户库失败（继续删除）：{type(ex).__name__}: {ex}')
        return ''


def db_delete_customer(website=''):
    """按 website_key 删掉一条客户，返回删除行数（0 或 1）。
    ⚠️ key 为空时**直接返回 0** —— normalize_domain 对坏 url 会返回 ''，
       少了这道门就会退化成 `DELETE FROM customers`（清库）。SQL 全参数化。
    ⚠️ 刻意不做 lower(email) 兜底：db_upsert_customer 会用 email 把新站**合并到别人那条**上，
       若这里再按 email 删，删掉的是另一家公司的记录。key 没命中就如实返回 0。"""
    key = normalize_domain(website)
    if not key:
        return 0
    with closing(db_connect()) as conn:
        row = conn.execute('SELECT id FROM customers WHERE website_key=? LIMIT 1',
                           (key,)).fetchone()
        if row is None:
            return 0
        cur = conn.execute('DELETE FROM customers WHERE id=?', (row['id'],))
        conn.commit()
        return cur.rowcount


QUERY_DELETE_MAX = 500      # 单次批量删除的硬上限（也顺带把 SQL 变量数压在 sqlite 限额内）


def db_delete_customers_by_ids(ids):
    """按**主键 id** 批量删客户，返回 (实际删除条数, 库里没找到的 id 列表)。

    与 db_delete_customer 的分工：
      · `db_delete_customer(website)` 按 website_key 删 —— 「本轮明细」手里只有 url，只能这么删；
      · 本函数按主键删 —— 查询页每行都带着 id，主键没有歧义（不存在键错位删错人、
        也不会像 email 兜底那样删到「合并进来的另一家」）。
    两条路都**不做 email 兜底**，口径一致。

    ⚠️ 安全门：ids 里非正整数一律丢弃；筛完为空就返回 (0, [])，**绝不能**拼出
       `DELETE FROM customers`（少了这道门，前端传个 []/None/字符串就等于清库）。
    """
    clean = []
    for v in (ids if isinstance(ids, (list, tuple)) else []):
        if isinstance(v, bool):
            continue                       # True/False 在 Python 里是 int，别把它当 id
        try:
            n = int(v)
        except (TypeError, ValueError):
            continue
        if n > 0:
            clean.append(n)
    clean = list(dict.fromkeys(clean))     # 去重：同一 id 传两次不该被算两条
    if not clean:
        return 0, []
    holes = ','.join('?' * len(clean))
    with closing(db_connect()) as conn:
        found = {r['id'] for r in conn.execute(
            'SELECT id FROM customers WHERE id IN (%s)' % holes, clean)}
        cur = conn.execute('DELETE FROM customers WHERE id IN (%s)' % holes, clean)
        conn.commit()
    return cur.rowcount, [i for i in clean if i not in found]


# ================= 客户库查询（只读，供「客户库查询页 /db」用）=================
# 这是一条**纯读**链路，两条硬纪律：
#   ① 不碰 STATE / STATE_LOCK —— 那把锁是非可重入的全局锁，一次几万行的扫描若占着它，
#      采集、落盘、页面轮询会一起卡住。查询只开自己的 sqlite 连接，出函数就关。
#   ② 不改 paths.ENRICH_FILE / STATE['output'] —— 那两个是「本轮结果文件」指针，
#      被查询串改的话，飞书导出会导到错的表上。
# 注入面：条件值全部走参数位；排序字段走白名单（键→子句），绝不把外部字符串拼进 SQL。

# 表格与导出要读的列（刻意不含 website_key：那是内部去重键，页面与业务上都用不到）
QUERY_SELECT_COLS = ('id, company_name, website, email, phone, contact_url, country, industry, '
                     'score, rating, snippet, search_formula, source, status, notes, '
                     'first_seen_at, last_seen_at')

# 排序白名单。score 为 NULL（还没评过分）的行一律沉到末尾，
# 否则「按评分降序」会把那几条没分的顶在最前面。
QUERY_ORDER_SQL = {
    'score_desc': 'score IS NULL, score DESC, id DESC',
    'score_asc':  'score IS NULL, score ASC, id DESC',
    'newest':     'first_seen_at DESC, id DESC',
    'oldest':     'first_seen_at ASC, id ASC',
    'updated':    'last_seen_at DESC, id DESC',
    'name':       'company_name COLLATE NOCASE ASC, id ASC',
}
QUERY_DEFAULT_ORDER = 'score_desc'
QUERY_MAX_LIMIT = 500          # 单页硬上限：防前端传巨大的 limit 把内存撑爆
QUERY_EMPTY = '__empty__'      # 多选里代表「这一列为空」的哨兵值（前端「未标注 / 未填写」用）

# 「联系方式」单选：值 → SQL 片段。列名全部写死在这里，不接受任何外部输入。
_QUERY_CONTACT_SQL = {
    'has_email':       "IFNULL(email,'')<>''",
    'has_phone':       "IFNULL(phone,'')<>''",
    'has_contact_url': "IFNULL(contact_url,'')<>''",
    'has_both':        "IFNULL(email,'')<>'' AND IFNULL(phone,'')<>''",
    'has_any':         "(IFNULL(email,'')<>'' OR IFNULL(phone,'')<>'')",
    'no_email':        "IFNULL(email,'')=''",
    'no_phone':        "IFNULL(phone,'')=''",
    'no_contact':      "IFNULL(email,'')='' AND IFNULL(phone,'')=''",
}

_DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')


def _query_where(f):
    """把筛选 dict 翻成 (WHERE 片段列表, 参数列表)。列名全是字面量，值全进参数位。"""
    where, args = [], []

    kw = str(f.get('kw') or '').strip()
    if kw:
        like = '%' + kw + '%'
        where.append('(company_name LIKE ? OR website LIKE ? OR email LIKE ? OR phone LIKE ?'
                     ' OR snippet LIKE ? OR search_formula LIKE ? OR industry LIKE ?'
                     ' OR country LIKE ? OR notes LIKE ?)')
        args += [like] * 9

    def _multi(col, values):
        """行业/国家/来源/评级这类多选：命中任一值即可；勾了「未填写」时空值也算命中。"""
        raw = [str(v) for v in (values or [])]
        picked = [v for v in raw if v != QUERY_EMPTY]
        parts = []
        if picked:
            parts.append(f'{col} IN ({",".join("?" * len(picked))})')
            args.extend(picked)
        if QUERY_EMPTY in raw:
            parts.append(f"IFNULL({col},'')=''")
        if parts:
            where.append('(' + ' OR '.join(parts) + ')')

    _multi('industry', f.get('industries'))
    _multi('country', f.get('countries'))
    _multi('source', f.get('sources'))
    _multi('rating', f.get('ratings'))

    for key in ('score_min', 'score_max'):
        v = f.get(key)
        if v in (None, ''):
            continue
        try:
            n = int(v)
        except (TypeError, ValueError):
            raise ValueError(f'{key} 必须是整数：{v!r}')
        where.append('score IS NOT NULL AND score ' + ('>=' if key == 'score_min' else '<=') + ' ?')
        args.append(n)

    rule = _QUERY_CONTACT_SQL.get(str(f.get('contact') or ''))
    if rule:
        where.append(rule)

    for key, op in (('seen_after', '>='), ('seen_before', '<=')):
        v = str(f.get(key) or '').strip()
        if not v:
            continue
        if not _DATE_RE.match(v):
            raise ValueError(f'{key} 必须是 YYYY-MM-DD 格式：{v!r}')
        # 只给到「天」，末端补足时分秒，否则「到 X 日为止」会把 X 日当天的记录全漏掉
        where.append(f'first_seen_at {op} ?')
        args.append(v + (' 00:00:00' if op == '>=' else ' 23:59:59'))
    return where, args


def db_query_customers(filters=None, order=QUERY_DEFAULT_ORDER, limit=50, offset=0):
    """按条件分页查询客户，返回 (命中总数, 当前页行列表, 实际生效的 limit)。

    计数与取行共用**同一个 WHERE**，所以「命中 N 条」与翻到底看到的一定自洽。
    order 只认 QUERY_ORDER_SQL 里的键；非法值静默回落默认排序（既不报错也不拼进 SQL）。
    """
    init_customer_db()
    where, args = _query_where(filters or {})
    clause = (' WHERE ' + ' AND '.join(where)) if where else ''
    order_sql = QUERY_ORDER_SQL.get(str(order or '')) or QUERY_ORDER_SQL[QUERY_DEFAULT_ORDER]
    try:
        limit = max(1, min(int(limit), QUERY_MAX_LIMIT))
    except (TypeError, ValueError):
        limit = 50
    try:
        offset = max(0, int(offset))
    except (TypeError, ValueError):
        offset = 0
    with closing(db_connect()) as conn:
        total = conn.execute('SELECT COUNT(*) FROM customers' + clause, args).fetchone()[0]
        rows = [dict(r) for r in conn.execute(
            f'SELECT {QUERY_SELECT_COLS} FROM customers{clause}'
            f' ORDER BY {order_sql} LIMIT ? OFFSET ?', args + [limit, offset])]
    return total, rows, limit


def db_facet_counts():
    """给筛选面板提供「库里真实存在的取值 + 各自条数」（含「未填写」项的条数）。

    ⚠️ 必须读真实 distinct 值，**不能**拿关键词库那 8 个规范行业名硬编码 ——
       库里还留着历史行业名（家具家居 18 条 / 汽配 5 条），硬编码会让这些数据一筛就漏。
    ⚠️ 下面四个列名都是本函数里写死的字面量，不接受任何外部输入
       （f-string 拼进去的全是常量，不是用户数据，无注入面）。
    """
    init_customer_db()
    facets = {}
    with closing(db_connect()) as conn:
        for key, col in (('industries', 'industry'), ('countries', 'country'),
                         ('sources', 'source'), ('ratings', 'rating')):
            facets[key] = [{'value': r[0], 'count': r[1]} for r in conn.execute(
                f"SELECT {col}, COUNT(*) FROM customers WHERE IFNULL({col},'')<>''"
                f' GROUP BY {col} ORDER BY COUNT(*) DESC, {col}')]
            facets[key + '_blank'] = conn.execute(
                f"SELECT COUNT(*) FROM customers WHERE IFNULL({col},'')=''").fetchone()[0]

        def _n(sql):
            return conn.execute(sql).fetchone()[0]

        facets['total'] = _n('SELECT COUNT(*) FROM customers')
        facets['contact'] = {
            'has_email': _n("SELECT COUNT(*) FROM customers WHERE IFNULL(email,'')<>''"),
            'has_phone': _n("SELECT COUNT(*) FROM customers WHERE IFNULL(phone,'')<>''"),
            'has_contact_url': _n("SELECT COUNT(*) FROM customers WHERE IFNULL(contact_url,'')<>''"),
            'no_contact': _n("SELECT COUNT(*) FROM customers "
                             "WHERE IFNULL(email,'')='' AND IFNULL(phone,'')=''"),
        }
        smin, smax = conn.execute(
            'SELECT MIN(score), MAX(score) FROM customers WHERE score IS NOT NULL').fetchone()
        facets['score'] = {'min': smin, 'max': smax}
        dmin, dmax = conn.execute(
            'SELECT MIN(substr(first_seen_at,1,10)), MAX(substr(first_seen_at,1,10))'
            ' FROM customers').fetchone()
        facets['date'] = {'min': dmin or '', 'max': dmax or ''}
    return facets
