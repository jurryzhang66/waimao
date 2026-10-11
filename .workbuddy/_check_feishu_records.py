# -*- coding: utf-8 -*-
"""build_feishu_records() 的行为断言 —— 本轮新加的两处逻辑必须有测试盯着。

被改动的两处：
  ① 「来源渠道」原来写死 `['谷歌搜索']`；现在按网址回客户库查真实 source
     （结果 Excel 的 10 列里没有来源列，契约不能改，所以只能回库查）。
  ② 50 分门槛对 `[手动采集]` 行**豁免** —— 那道门槛是拦自动搜索捞进来的垃圾站的，
     手动采集是使用者自己认定「这就是我的目标客户」，不该被静默丢掉。

数据库与结果文件都指向临时目录，绝不触碰真实客户库。

用法：python .workbuddy\\_check_feishu_records.py      （EXIT 0 = 全过）
"""
import importlib.util
import os
import shutil
import sqlite3
import sys
import tempfile

sys.stdout.reconfigure(encoding='utf-8')

SRC = r'D:\waimao\dp\客户开发工具.py'
spec = importlib.util.spec_from_file_location('kf_fs', SRC)
m = importlib.util.module_from_spec(spec)
sys.modules['kf_fs'] = m
spec.loader.exec_module(m)

TMP = tempfile.mkdtemp(prefix='check_feishu_')
OUT = os.path.join(TMP, 'output')
os.makedirs(OUT, exist_ok=True)
m.db_path = lambda: os.path.join(TMP, 'test.sqlite3')
m.db.db_path = m.db_path
m.paths.ENRICH_FILE = os.path.join(OUT, 'result.xlsx')
m.ts_result_path = lambda industry='', country='': m.paths.ENRICH_FILE
m.STATE['output'] = ''
assert os.path.commonpath([os.path.realpath(m.db.db_path()), os.path.realpath(TMP)]) == os.path.realpath(TMP)
assert os.path.commonpath([os.path.realpath(m.paths.ENRICH_FILE), os.path.realpath(TMP)]) == os.path.realpath(TMP)
m.init_customer_db()

FAILS = []


def check(cond, label, detail=''):
    print(('  ✅ ' if cond else '  ❌ ') + label + (f' —— {detail}' if detail and not cond else ''))
    if not cond:
        FAILS.append(label)


def row(title, url, score, formula, rating='观察', country='', emails=(), industry='塑料包装袋'):
    return {'title': title, 'url': url, 'emails': list(emails), 'phones': [],
            'contact_url': '', 'score': score, 'rating': rating, 'snippet': '摘要',
            'formula': formula, 'country': country, 'industry': industry}


ROWS = [
    # 手动采集 + 低分 → 必须豁免 50 分门槛导出
    row('Manual Low Score Co', 'https://manual-low.com', 30,
        '[手动采集] https://manual-low.com', emails=['info@manual-low.com']),
    # 自动链路 + 低分 → 必须被门槛拦掉
    row('Auto Low Score Co', 'https://auto-low.com', 30, 'cat litter wholesaler',
        emails=['info@auto-low.com'], country='法国'),
    # 自动链路 + 高分 → 正常导出
    row('Auto High Co', 'https://auto-high.com', 75, 'cat litter distributor',
        rating='重点客户', emails=['info@auto-high.com'], country='法国'),
    # 手动采集 + 无评分 → 不能因为 int(None) 崩掉
    # （score 先给 0 占位，落盘后手工把那一格清空 —— _save_enrich_excel 本身按 score 排序，
    #   塞 None 会直接 TypeError；而「空评分」正是启动迁移老结果时真会出现的格子）
    row('Manual No Score Co', 'https://manual-noscore.com', 0,
        '[手动采集] https://manual-noscore.com'),
    # 手动采集 + 国家需要映射（沙特 → 沙特阿拉伯）
    row('Manual Saudi Co', 'https://manual-saudi.com', 55,
        '[手动采集] https://manual-saudi.com', country='沙特', emails=['info@manual-saudi.com']),
    # 客户库里来源是个飞书表里没有的怪值 → 应回落「谷歌搜索」并记警告，不能整条崩
    row('Weird Source Co', 'https://weird-source.com', 60, 'cat litter supplier'),
    # 多邮箱（新格式：逗号）→「邮箱」取第一个，其余写进备注，且用逗号分隔（用户要求）
    row('Multi Comma Co', 'https://multi-comma.com', 70, 'cat litter supplier',
        emails=['first@multi-comma.com', 'second@multi-comma.com'], country='法国'),
    # 多邮箱（**旧格式**：斜杠）→ 由下面手工改写 Excel 邮箱格制造，验证 split_emails 向后兼容
    row('Multi Slash Co', 'https://multi-slash.com', 70, 'cat litter supplier',
        emails=['one@multi-slash.com', 'two@multi-slash.com'], country='法国'),
]
m._save_enrich_excel(ROWS, industry='塑料包装袋')
check(os.path.exists(m.paths.ENRICH_FILE), '临时结果文件已生成')

# 把 Manual No Score 那一行的「评分」格清空，制造空评分
from openpyxl import load_workbook            # noqa: E402
_wb = load_workbook(m.paths.ENRICH_FILE)
_ws = _wb.active
_blanked = 0
for _r in range(2, _ws.max_row + 1):
    if _ws.cell(_r, 2).value == 'https://manual-noscore.com':
        _ws.cell(_r, 6).value = None
        _blanked += 1
    elif _ws.cell(_r, 2).value == 'https://multi-slash.com':
        # 制造**旧格式**邮箱格（2026-10-11 之前 Excel 里写的是 ' / '）——
        # 用来验证飞书导出的 split_emails 能向后兼容存量数据（只认逗号会把这行邮箱全丢）。
        _ws.cell(_r, 3).value = 'one@multi-slash.com / two@multi-slash.com'
_wb.save(m.paths.ENRICH_FILE)
_wb.close()
check(_blanked == 1, '把「Manual No Score Co」那行的评分格清空了（制造空评分样本）', f'blanked={_blanked}')

# 往临时库里写来源（build_feishu_records 就是按 website_key 回库查这张表）
conn = sqlite3.connect(m.db_path())
conn.executemany(
    "INSERT INTO customers (company_name, website, website_key, source) VALUES (?,?,?,?)",
    [('Manual Low', 'https://manual-low.com', 'manual-low.com', '手动采集'),
     ('Manual Saudi', 'https://manual-saudi.com', 'manual-saudi.com', '手动采集'),
     ('Auto High', 'https://auto-high.com', 'auto-high.com', '谷歌地图'),
     ('Weird', 'https://weird-source.com', 'weird-source.com', '微博')])
conn.commit()
conn.close()

print('\n========== 生成飞书记录 ==========')
recs = m.build_feishu_records()
by_url = {r['官网']: r for r in recs}
for r in recs:
    print('   ', r['公司名称'], '|', r['来源渠道'], '|', r.get('国家/地区'), '|', r['备注'])

print('\n========== ① 50 分门槛与豁免 ==========')
check('https://manual-low.com' in by_url, '手动采集 + 30 分 → **导出**（豁免门槛）')
check('https://auto-low.com' not in by_url, '自动链路 + 30 分 → **不导出**（门槛仍然生效）')
check('https://auto-high.com' in by_url, '自动链路 + 75 分 → 导出')
check(len(recs) == 7, f'共 7 条（8 行里只砍掉 auto-low）（实际 {len(recs)}）')

print('\n========== ② 来源渠道按网址回库查 ==========')
check(by_url.get('https://manual-low.com', {}).get('来源渠道') == ['手动采集'],
      '库 source=手动采集 → 飞书写「手动采集」',
      str(by_url.get('https://manual-low.com', {}).get('来源渠道')))
check(by_url.get('https://auto-high.com', {}).get('来源渠道') == ['谷歌地图'],
      '库 source=谷歌地图 → 飞书写「谷歌地图」',
      str(by_url.get('https://auto-high.com', {}).get('来源渠道')))
check(by_url.get('https://manual-noscore.com', {}).get('来源渠道') == ['谷歌搜索'],
      '库里查不到该网址 → 回落「谷歌搜索」',
      str(by_url.get('https://manual-noscore.com', {}).get('来源渠道')))
check(by_url.get('https://weird-source.com', {}).get('来源渠道') == ['谷歌搜索'],
      '库里的怪取值（微博）→ 回落「谷歌搜索」而不是把整条写坏',
      str(by_url.get('https://weird-source.com', {}).get('来源渠道')))
check(all(isinstance(r['来源渠道'], list) and len(r['来源渠道']) == 1 for r in recs),
      '来源渠道始终是**单元素数组**（飞书 select 的写法）')

print('\n========== ③ 无评分的手动行不能崩 ==========')
ns = by_url.get('https://manual-noscore.com')
check(ns is not None, 'score 为 None 的手动行仍在（没被 int(None) 炸掉）')
check(ns and '0分' in ns['备注'], '备注里评分按 0 显示', ns and ns['备注'])
check(ns and '手动采集' in ns['备注'], '手动采集会在备注里留痕', ns and ns['备注'])

print('\n========== ④ 其余字段仍受 10 列契约约束 ==========')
check(by_url.get('https://manual-low.com', {}).get('邮箱') == 'info@manual-low.com',
      '邮箱取第一个合规邮箱')
check('其他邮箱' not in by_url.get('https://manual-low.com', {}).get('备注', ''),
      '只有一个邮箱时不写「其他邮箱」')
check(by_url.get('https://manual-saudi.com', {}).get('国家/地区')
      == [m.COUNTRY_STANDARD.get('沙特', '沙特')]
      and by_url.get('https://manual-saudi.com', {}).get('国家/地区') == ['沙特阿拉伯'],
      '国家做了标准映射（沙特 → 沙特阿拉伯）',
      str(by_url.get('https://manual-saudi.com', {}).get('国家/地区')))
check(by_url.get('https://manual-low.com', {}).get('跟进状态') == ['未联系'],
      '跟进状态默认「未联系」')
check(by_url.get('https://manual-low.com', {}).get('使用的搜索公式') == '[手动采集] https://manual-low.com',
      '搜索公式原样带过去（飞书那边靠它追溯来源）')

print('\n========== ⑤ 多邮箱：逗号分隔 + 兼容旧斜杠数据 ==========')
mc = by_url.get('https://multi-comma.com', {})
check(mc.get('邮箱') == 'first@multi-comma.com',
      '多邮箱只取第一个进飞书「邮箱」字段（该字段只收单个邮箱）', str(mc.get('邮箱')))
check('其他邮箱：second@multi-comma.com' in (mc.get('备注') or ''),
      '其余邮箱进备注「其他邮箱」，且用逗号分隔（不是斜杠）', str(mc.get('备注')))
check(' / ' not in (mc.get('备注') or ''), '逗号格式的行，备注里不该再出现斜杠', str(mc.get('备注')))
ms = by_url.get('https://multi-slash.com', {})
check(ms.get('邮箱') == 'one@multi-slash.com',
      '**旧斜杠格式**的存量数据仍能拆出第一个邮箱（split_emails 向后兼容）', str(ms.get('邮箱')))
check('其他邮箱：two@multi-slash.com' in (ms.get('备注') or ''),
      '旧格式的其余邮箱也照旧进备注（没被当成一个长邮箱静默丢掉）', str(ms.get('备注')))

print('\n' + '=' * 62)
shutil.rmtree(TMP, ignore_errors=True)
if FAILS:
    print(f'❌ 失败 {len(FAILS)} 项：')
    for f in FAILS:
        print('   · ' + f)
    sys.exit(1)
print('✅ build_feishu_records 的来源渠道与门槛豁免行为全部正确')
