# -*- coding: utf-8 -*-
"""打桩冒烟：手动采集的「本轮会话」—— 实时明细 + 结束时一次性落盘。

对应用户需求：「采集我怎么知道有没有成功？」「结束的时候要生成 Excel，像自动采集那样」。
本脚本把浏览器/网络换成假实现，让真实的 collect_from_html / _flush_manual_session /
两个会话路由跑起来；数据库与 output 都指向临时目录，绝不碰真实客户库。

覆盖 12 组：
  ① 无会话时扩展单点采集 → 自动开一轮（口径取请求里的行业/国家），且**不写库、不落盘**
  ② /api/manual/batch 并入同一轮 → 批量子过程字段与本轮累计字段各管各的
  ③ 结束时**先入库**、再落一份 10 列结果文件（降序、成为「最新结果文件」、本轮清零）
  ④ 文件名口径 = 开轮那一刻的行业/国家；同分钟再落一次自动加 _2，不覆盖
  ⑤ 空轮直接结束 → 不产文件、不动「最新结果文件」指针
  ⑥ 本轮还有内容时再点「开始本轮」→ 被拒，且不清空已有明细
  ⑦ 批量还在跑时点「结束」→ 立即返回 pending，由 worker 收尾落盘（不阻塞、不丢行）
  ⑧ 落盘失败（Excel 被占用）→ 保留本轮数据 + 错误提示；恢复后一键重试成功
  ⑨ /api/manual/state 的 session_* 字段契约
  ⑩ 自动任务运行时拒绝开本轮
  ⑪ 删除单条明细：本轮还没入库 → 只动内存（不碰库/不产备份）；附老客户灰字行（绝不动库）
     与「已入库但 Excel 没写成功」（唯一会动库、需双删）两条分支
  ⑫ 去重键自愈：website_key 恒等于 normalize_domain(website)（否则同站会被反复算「新客户」）

用法：python .workbuddy\\_smoke_manual_session.py
"""
import importlib.util
import os
import shutil
import sqlite3
import sys
import tempfile
import time

from openpyxl import load_workbook

SRC = r'D:\waimao\dp\客户开发工具.py'
spec = importlib.util.spec_from_file_location('kf', SRC)
m = importlib.util.module_from_spec(spec)
sys.modules['kf'] = m
spec.loader.exec_module(m)

sys.stdout.reconfigure(encoding='utf-8')

_ok, _fail = 0, []


def check(cond, desc, extra=''):
    global _ok
    if cond:
        _ok += 1
        print(f'  ✅ {desc}')
    else:
        _fail.append(desc)
        print(f'  ❌ {desc}' + (f'   ← {extra}' if extra else ''))


def section(title):
    print(f'\n================ {title} ================')


TMP = tempfile.mkdtemp(prefix='smoke_session_')
OUT = os.path.join(TMP, 'output')
os.makedirs(OUT, exist_ok=True)

# ---- 护栏：库与输出全部关进临时目录（漏一个就会写脏真实数据）----
m.db_path = lambda: os.path.join(TMP, 'test.sqlite3')
m.db.db_path = m.db_path
m.paths.output_dir = lambda: OUT                     # ts_result_path / find_latest_result 都走它
m.paths.ENRICH_FILE = os.path.join(OUT, 'main.xlsx')


def _under_tmp(p):
    return os.path.commonpath([os.path.realpath(p),
                               os.path.realpath(TMP)]) == os.path.realpath(TMP)


check(_under_tmp(m.db.db_path()), 'core.db.db_path 已指向临时目录')
check(_under_tmp(m.paths.output_dir()), 'paths.output_dir() 已指向临时目录')
check(_under_tmp(m.paths.ENRICH_FILE), 'ENRICH_FILE 已指向临时目录')
check(_under_tmp(m.ts_result_path('塑料包装袋', '德国')), '真实 ts_result_path 落在临时目录')
# ⚠️ 删除单条会「删前自动备份客户库」→ 备份目录由 db_path() 推导，必须一起关进临时目录，
#    否则 ⑪ 组会往真实 DB/backups/ 里写测试备份文件。
check(_under_tmp(m.db._backups_dir()), '删除前备份目录也落在临时目录（由 db_path 推导）')
print('打桩护栏 OK：', TMP)

# ---- 网络/浏览器打桩：秒级 HTML 命中，或升级到假 enrich ----
calls = {'enrich': [], 'req': 0}


class FakePage:
    def quit(self):
        pass


m.create_browser_page = lambda *a, **k: FakePage()
m._manual_browser = lambda: FakePage()


def fake_enrich(page, url):
    calls['enrich'].append(url)
    host = m.normalize_domain(url)
    return {'emails': [f'sales@{host}'], 'phones': [],
            'contact_url': url + '/contact', 'china_signal_text': ''}


m.enrich = fake_enrich


class FakeRequests:
    @staticmethod
    def get(*a, **k):
        calls['req'] += 1
        raise OSError('stubbed network failure')


m.requests = FakeRequests()


def reset_state():
    m.MANUAL_STATE.update({'running': False, 'stop': False, 'total': 0, 'done': 0,
                           'added': 0, 'existed': 0, 'failed': 0, 'upgraded': 0,
                           'current': '', 'ext_last_seen': 0.0,
                           'session_active': False, 'session_started_at': 0.0,
                           'session_rows': [], 'session_rid_seq': 0,
                           'session_industry': '', 'session_country': '',
                           'session_excel': '', 'session_error': '',
                           'session_finish_pending': False})
    m.STATE['running'] = False


reset_state()


def state():
    with m.STATE_LOCK:
        st = dict(m.MANUAL_STATE)
        st['session_rows'] = list(m.MANUAL_STATE['session_rows'])
    return st


def files():
    return sorted(os.listdir(OUT))


def wait_until(pred, timeout=25, interval=0.15):
    t0 = time.time()
    while time.time() - t0 < timeout:
        if pred():
            return True
        time.sleep(interval)
    return False


def excel_rows(path):
    wb = load_workbook(path)
    ws = wb.active
    head = [c.value for c in ws[1]]
    data = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    wb.close()
    return head, data


def html_with_mail(host):
    return (f'<html><head><title>Trash Bags Wholesale Supplier</title>'
            f'<meta name="description" content="Leading wholesale trash bag supplier"></head>'
            f'<body><a href="mailto:info@{host}">info@{host}</a></body></html>')


section('① 无会话时扩展单点 → 自动开一轮（且不落盘）')
with m.app.test_client() as c:
    r = c.post('/api/collect',
               json={'url': 'https://alpha-pack.de', 'title': 'Alpha Pack GmbH',
                     'html': html_with_mail('alpha-pack.de'),
                     'industry': '塑料包装袋', 'country': '德国'},
               headers={'X-DP-Ext': '1'})
    j = r.get_json() or {}
check(j.get('ok') and j.get('added'), '扩展单点采集成功', str(j))
st = state()
check(st['session_active'] is True, '本轮被自动开起来了（用户不必先切回页面点开始）')
check(st['session_industry'] == '塑料包装袋', '本轮行业口径 = 请求里的行业', st['session_industry'])
check(st['session_country'] == '德国', '本轮国家口径 = 请求里的国家', st['session_country'])
check(len(st['session_rows']) == 1, '明细里记上了这 1 条', str(len(st['session_rows'])))
check(st['session_rows'][0]['kind'] == 'added', '明细行带 kind=added')
check(st['session_excel'] == '', '还没结束，session_excel 应为空')
check(m.paths.ENRICH_FILE == os.path.join(OUT, 'main.xlsx'), '采集阶段没碰「最新结果文件」')
check(files() == [], f'采集阶段不该产出任何 xlsx，实际 {files()}')
# ⛔ 本轮改动的核心：采集阶段**也不写客户库** —— 入库整轮只发生一次，就在「结束」那一刻。
check(not m.db_has_customer('https://alpha-pack.de'),
      '① 采集阶段不写客户库（入库时机已移到「结束并生成 Excel」）')
check('kind' not in (j.get('row') or {}), '返回给扩展的 row 不带内部 kind 字段')

section('② /api/manual/batch 并入同一轮（本轮口径不被批量表单覆盖）')
with m.app.test_client() as c:
    rb = c.post('/api/manual/batch', json={
        'urls': 'https://beta-liner.co.uk\nhttps://gamma-bag.fr',
        'industry': 'LED照明', 'country': '法国',   # ← 与①不同，应被忽略（本轮口径只认第一次）
        'upgrade': True,
    })
    jb = rb.get_json() or {}
check(jb.get('ok') and jb.get('total') == 2, '批量已受理 2 条', str(jb))
ok = wait_until(lambda: state()['done'] == 2 and not state()['running'])
check(ok, '批量跑完（done=2、running=False）', str(state()))
st = state()
check(st['session_industry'] == '塑料包装袋', '本轮行业口径没被批量表单覆盖', st['session_industry'])
check(st['session_country'] == '德国', '本轮国家口径没被批量表单覆盖', st['session_country'])
check(st['total'] == 2 and st['done'] == 2, '批量子过程字段是「本次批量」语义')
check(st['added'] == 3, f'本轮累计新增 3 条（1 单点 + 2 批量），实际 {st["added"]}')
check(len(st['session_rows']) == 3, '明细累计 3 条', str(len(st['session_rows'])))
check(files() == [], '批量跑完也还没落盘（要等结束）')

section('③ 结束本轮 → 先入库，再落一份 10 列结果文件')
with m.app.test_client() as c:
    rf = c.post('/api/manual/session/finish')
    jf = rf.get_json() or {}
check(jf.get('ok') and jf.get('saved'), '结束接口返回已保存', str(jf))
check(jf.get('rows') == 3, '返回本轮 3 条', str(jf.get('rows')))
new_files = files()
check(len(new_files) == 1, f'产出且只产出一份文件，实际 {new_files}')
saved = m.paths.ENRICH_FILE
check(os.path.dirname(saved) == OUT, '文件落在 output 目录', saved)
name = os.path.basename(saved)
check(name.startswith('塑料包装袋-') and '-德国.xlsx' in name,
      '文件名 = 行业（开轮口径）-时间-国家', name)
check('output/' in (jf.get('file') or ''), 'file 是 output/ 相对串', str(jf.get('file')))
check(m.STATE['output'].startswith('output/'), 'STATE["output"] 也更新了', m.STATE['output'])
head, data = excel_rows(saved)
check(head == m.EXCEL_HEADER_10, '10 列硬契约', str(head))
check(len(data) == 3, f'3 行数据，实际 {len(data)}')
scores = [r[5] for r in data]
check(scores == sorted(scores, reverse=True), f'按评分降序：{scores}')
check(all(str(r[8]).startswith('[手动采集]') for r in data), '公式列带 [手动采集] 前缀')
st = state()
check(st['session_active'] is False, '结束后本轮关闭')
check(st['session_rows'] == [], '结束后明细清空')
check(st['session_excel'] == name, 'session_excel 记下了文件名', st['session_excel'])
check(not st['session_error'], '结束后没有残留错误')
# ⛔ 入库时机：整轮只在**点「结束」这一刻**写客户库（采集时一行都不写）。
check(m.db_has_customer('https://alpha-pack.de'), '③ 结束时才入库：单点那条已在客户库')
check(m.db_has_customer('https://beta-liner.co.uk') and m.db_has_customer('https://gamma-bag.fr'),
      '③ 批量那 2 条也在结束这一刻入库了')

section('④ 文件名口径只认开轮那一刻；同分钟再落一次加 _2 不覆盖')
with m.app.test_client() as c:
    rs = c.post('/api/manual/session/start', json={'industry': '塑料包装袋', 'country': '德国'})
    js = rs.get_json() or {}
check(js.get('ok'), '显式开一轮成功', str(js))
with m.app.test_client() as c:
    c.post('/api/collect', json={'url': 'https://delta-pack.es', 'title': 'Delta Pack SL',
                                 'html': html_with_mail('delta-pack.es')},
           headers={'X-DP-Ext': '1'})
    rf2 = c.post('/api/manual/session/finish')
    jf2 = rf2.get_json() or {}
check(jf2.get('saved'), '第二次结束也成功', str(jf2))
name2 = os.path.basename(m.paths.ENRICH_FILE)
check(name2 != name, f'没有覆盖上一份（{name2} ≠ {name}）')
check('_2' in name2, f'同分钟同口径自动加 _2：{name2}')
check(len(files()) == 2, f'现在共 2 份结果文件，实际 {files()}')

section('⑤ 空轮直接结束 → 不产文件、不动指针')
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={'industry': '服装', 'country': '日本'})
    before_files, before_ptr = files(), m.paths.ENRICH_FILE
    re_ = c.post('/api/manual/session/finish')
    je = re_.get_json() or {}
check(je.get('ok') and je.get('saved') is False, '空轮结束返回 ok 但 saved=False', str(je))
check('没有采到新客户' in (je.get('msg') or ''), '给出「本轮没有采到新客户」的说明', str(je.get('msg')))
check(files() == before_files, '没产出空文件', str(files()))
check(m.paths.ENRICH_FILE == before_ptr, '「最新结果文件」指针没动')
check(state()['session_active'] is False, '空轮也把本轮关掉了')

section('⑥ 本轮还有内容时再点「开始本轮」→ 被拒且不清空')
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={})
    c.post('/api/collect', json={'url': 'https://epsilon-pack.it', 'title': 'Epsilon Pack Srl',
                                 'html': html_with_mail('epsilon-pack.it')},
           headers={'X-DP-Ext': '1'})
    n_before = len(state()['session_rows'])
    r2 = c.post('/api/manual/session/start', json={})
    j2 = r2.get_json() or {}
check(j2.get('ok') is False, '第二次开轮被拒', str(j2))
check('结束' in (j2.get('msg') or ''), '提示里让用户先结束', str(j2.get('msg')))
check(len(state()['session_rows']) == n_before == 1, '已有明细没被清空', str(n_before))

section('⑦ 批量还在跑时点「结束」→ pending，由 worker 收尾落盘')
with m.app.test_client() as c:
    c.post('/api/manual/batch', json={
        'urls': 'https://zeta-bag.pl\nhttps://eta-liner.se\nhttps://theta-pack.pt',
        'upgrade': True})
    time.sleep(0.25)                      # 让它至少跑起来（不阻塞、也不等它跑完）
    rp = c.post('/api/manual/session/finish')
    jp = rp.get_json() or {}
check(jp.get('ok') and jp.get('pending') is True, '批量进行中 → 立即返回 pending', str(jp))
check(state()['stop'] is True, '同时置了批量停止标志')
ok = wait_until(lambda: not state()['running'] and not state()['session_active'], timeout=40)
check(ok, 'worker 收尾后本轮已关闭', str(state()))
check(m.paths.ENRICH_FILE != '', '收尾后指针有效')
head, data = excel_rows(m.paths.ENRICH_FILE)
check(len(data) >= 2, f'文件含已采到的行（1 单点 + 批量里的若干），实际 {len(data)}')
check(all(str(r[8]).startswith('[手动采集]') for r in data), '收尾这份也是 10 列契约的数据')
check(state()['session_rows'] == [], '收尾后明细清空')

section('⑧ 落盘失败（Excel 被占用）→ 保留数据，恢复后重试成功')
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={'industry': '家具', 'country': '越南'})
    c.post('/api/collect', json={'url': 'https://iota-furniture.vn', 'title': 'Iota Furniture',
                                 'html': html_with_mail('iota-furniture.vn')},
           headers={'X-DP-Ext': '1'})
    _real = m._save_enrich_excel
    m._save_enrich_excel = lambda *a, **k: False
    before_files, before_ptr = files(), m.paths.ENRICH_FILE
    rf3 = c.post('/api/manual/session/finish')
    jf3 = rf3.get_json() or {}
    m._save_enrich_excel = _real
    check(jf3.get('ok') is False, '失败时如实返回 ok=False', str(jf3))
    check((jf3.get('msg') or '').strip(), '给出可读的失败原因', str(jf3.get('msg')))
    st = state()
    check(st['session_active'] is True, '失败后本轮保持进行中（好让用户重试）')
    check(len(st['session_rows']) == 1, '失败后明细保留（数据没白采）', str(len(st['session_rows'])))
    check(bool(st['session_error']), 'session_error 有内容')
    check(files() == before_files, '失败时没产出半截文件', str(files()))
    check(m.paths.ENRICH_FILE == before_ptr, '失败时指针没动')
    rr = c.post('/api/manual/session/finish')
    jr = rr.get_json() or {}
    check(jr.get('ok') and jr.get('saved'), '恢复后重试成功', str(jr))
    st = state()
    check(st['session_active'] is False and st['session_rows'] == [], '重试成功后本轮清零')
    check(not st['session_error'], '成功后错误信息被清掉')

section('⑨ /api/manual/state 的 session_* 字段契约')
with m.app.test_client() as c:
    st = c.get('/api/manual/state').get_json() or {}
for k in ('session_active', 'session_started_at', 'session_rows', 'session_industry',
          'session_country', 'session_excel', 'session_error',
          'added', 'existed', 'failed', 'upgraded', 'running', 'stop', 'total', 'done',
          'current', 'ext_last_seen', 'now'):
    check(k in st, f'/api/manual/state 含字段 {k}')
check(isinstance(st.get('session_rows'), list), 'session_rows 是列表')
check('recent' in st, '老的 recent 字段仍保留（不破坏既有消费方）')

section('⑩ 自动任务运行时拒绝开本轮（两条链路不互相踩）')
m.STATE['running'] = True
try:
    with m.app.test_client() as c:
        rw = c.post('/api/manual/session/start', json={})
        jw = rw.get_json() or {}
    check(jw.get('ok') is False, '自动搜索跑着时开本轮被拒', str(jw))
finally:
    m.STATE['running'] = False

section('⑪ 删除单条明细：本轮还没入库 → 只动内存；老客户/灰字行绝不连带删库')
# 用户需求：「采集到的有些数据不是自己想要的，想手动删掉」。
# ⚠️ 2026-10-11 起手动采集改成「点『🏁 结束并生成 Excel』才一次性入库」，所以本轮明细里
#    kind='added' 的行**此刻还不在客户库里** —— 删它只是把它从「待入库清单」里拿掉。
#    三条删除路径的行为已统一为：
#      · kind='added' 且**库里真的没有** → 只删内存明细（客户库不动、也不该平白备份）；
#      · kind='added' 且**库里已经有**（唯一来源：本轮已入库成功、但 Excel 没写成功，见 ⑪c）
#        → 明细 + 客户库双删（否则又是「假删除」）；
#      · kind='existed' 灰字行（别人早采过的老客户）→ **永远不碰库**（见 ⑪b，用户明确要求）。
with m.app.test_client() as c:
    rs0 = c.post('/api/manual/session/start', json={'industry': '食品', 'country': '意大利'})
    check((rs0.get_json() or {}).get('ok'), '⑪ 前置：开一轮成功', str(rs0.get_json()))
    for u, t in (('https://alpha-del.it', 'Alpha Del SpA'),
                 ('https://beta-keep.it', 'Beta Keep Srl'),
                 ('https://gamma-drop.it', 'Gamma Drop Snc')):
        c.post('/api/collect', json={'url': u, 'title': t, 'html': html_with_mail(u.rsplit('/', 1)[1])},
               headers={'X-DP-Ext': '1'})
st = state()
rids = [r['rid'] for r in st['session_rows']]
check(len(st['session_rows']) == 3, '⑪ 前置：明细 3 条', str(len(st['session_rows'])))
check(len(set(rids)) == 3 and all(isinstance(r, int) for r in rids),
      '⑪ 前置：rid 各不相同且是整数', str(rids))
check(st['session_rid_seq'] == 3, '⑪ rid 发号器走到 3', str(st['session_rid_seq']))
check(st['added'] == 3, '⑪ 前置：added 计数 = 3', str(st['added']))
check(not m.db_has_customer('https://alpha-del.it'), '⑪ 前置：本轮还没结束 → 库里**不该**有它')
check(not m.db_has_customer('https://beta-keep.it'), '⑪ 前置：保留站此刻也不在库里')

del_rid, keep_rid = rids[0], rids[1]
bk_before = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
with m.app.test_client() as c:
    jd = c.post('/api/manual/session/delete', json={'rid': del_rid}).get_json() or {}
bk_after = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
check(jd.get('ok') is True, '⑪ 删除返回 ok=True', str(jd))
check(jd.get('deleted') == 0, '⑪ 本轮未入库 → 客户库应删 0 条', str(jd.get('deleted')))
check(jd.get('rows') == 2, '⑪ 删完明细剩 2 条', str(jd.get('rows')))
check(jd.get('backup') == '', '⑪ 没动库就不该平白备份一份', str(jd.get('backup')))
check(bk_after == bk_before, '⑪ backups 目录不该新增文件', str(bk_after))
st = state()
check(len(st['session_rows']) == 2, '⑪ 明细由 3 条变 2 条', str(len(st['session_rows'])))
check(all(r['rid'] != del_rid for r in st['session_rows']), '⑪ 被删的那条已不在明细里')
check(any(r['rid'] == keep_rid for r in st['session_rows']), '⑪ 另一条没被连坐（按 rid 删，不是按 url）')
check('alpha-del' not in ' '.join(r.get('url', '') for r in st['session_rows']),
      '⑪ 明细里已搜不到被删域名')
check(st['added'] == 2, '⑪ added 计数回退到 2', str(st['added']))
check(not m.db_has_customer('https://alpha-del.it'), '⑪ 客户库本来就没有它，也不该被碰')
check(not m.db_has_customer('https://beta-keep.it'), '⑪ 保留站保持未入库状态')
with m.app.test_client() as c:
    st2 = c.get('/api/manual/state').get_json() or {}
check(len(st2.get('session_rows') or []) == 2, '⑪ /api/manual/state 反映新长度')
check(st2.get('added') == 2, '⑪ /api/manual/state 的 added 同步为 2', str(st2.get('added')))

# 删不存在 / 不带 rid → 一律拒绝（绝不能退化成「删全部」或「清库」）
with m.app.test_client() as c:
    jn = c.post('/api/manual/session/delete', json={'rid': 999999}).get_json() or {}
check(jn.get('ok') is False, '⑪ 删不存在的 rid → ok=False', str(jn))
with m.app.test_client() as c:
    jnn = c.post('/api/manual/session/delete', json={}).get_json() or {}
check(jnn.get('ok') is False, '⑪ 不带 rid 的请求也被拒', str(jnn))
check(len(state()['session_rows']) == 2, '⑪ 被拒的请求没动过明细')

with m.app.test_client() as c:
    jk = c.post('/api/manual/session/delete', json={'rid': keep_rid}).get_json() or {}
check(jk.get('ok') is True and jk.get('deleted') == 0, '⑪ 再删一条（同样未入库）也不碰库', str(jk))

# 结束落盘：被删的条目绝不能出现在 Excel 里，也不该被写进库
with m.app.test_client() as c:
    jfin = c.post('/api/manual/session/finish').get_json() or {}
check(jfin.get('ok') and jfin.get('saved'), '⑪ 剩余 1 条 → 结束落盘成功', str(jfin))
check(jfin.get('rows') == 1, '⑪ 落盘行数 = 剩余 1 条', str(jfin.get('rows')))
head, data = excel_rows(m.paths.ENRICH_FILE)
check(len(data) == 1, f'⑪ Excel 只有 1 行（被删行没进去），实际 {len(data)}')
urls_in_excel = ' '.join(str(r[1] or '') for r in data)
check('gamma-drop' in urls_in_excel, '⑪ Excel 里剩下的是没被删的那条', urls_in_excel)
check('alpha-del' not in urls_in_excel and 'beta-keep' not in urls_in_excel,
      '⑪ Excel 里不含被删的两条', urls_in_excel)
check(m.db_has_customer('https://gamma-drop.it'), '⑪ 结束这一刻剩下那条才入库')
check(not m.db_has_customer('https://alpha-del.it')
      and not m.db_has_customer('https://beta-keep.it'), '⑪ 被删的两条始终没进库')

section('⑪b 灰字行（已在库的老客户）删了绝不连带删库')
# 灰字行是**别人早先采的**老客户，这一轮只是「碰上了、跳过去」。删它不是要删人家的客户 ——
# 用户 2026-10-11 明确拍板「灰字行不再动库」。
m.db_upsert_customer({'company_name': 'Legacy Grey Ltd', 'website': 'https://legacy-grey.it',
                      'email': 'info@legacy-grey.it', 'source': '谷歌搜索'})
check(m.db_has_customer('https://legacy-grey.it'), '⑪b 前置：老客户已在库里')
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={'industry': '食品', 'country': '意大利'})
    c.post('/api/collect', json={'url': 'https://legacy-grey.it', 'title': 'Legacy Grey Ltd',
                                 'html': html_with_mail('legacy-grey.it')},
           headers={'X-DP-Ext': '1'})
    c.post('/api/collect', json={'url': 'https://fresh-one.it', 'title': 'Fresh One Srl',
                                 'html': html_with_mail('fresh-one.it')},
           headers={'X-DP-Ext': '1'})
st = state()
grey = [r for r in st['session_rows'] if r.get('kind') == 'existed']
check(len(grey) == 1 and grey[0]['url'] == 'https://legacy-grey.it',
      '⑪b 老客户被标成灰字（kind=existed）',
      str([(r.get('kind'), r.get('url')) for r in st['session_rows']]))
check(st['added'] == 1 and st['existed'] == 1, '⑪b 计数：有效 1 / 已在库 1',
      str({'added': st['added'], 'existed': st['existed']}))
bk_before2 = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
with m.app.test_client() as c:
    jg = c.post('/api/manual/session/delete', json={'rid': grey[0]['rid']}).get_json() or {}
bk_after2 = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
check(jg.get('ok') is True and jg.get('deleted') == 0, '⑪b 删灰字行 → 客户库删 0 条', str(jg))
check(jg.get('backup') == '' and bk_after2 == bk_before2, '⑪b 删灰字行不产生备份', str(jg))
check(m.db_has_customer('https://legacy-grey.it'),
      '⑪b ⛔ 老客户必须还在库里（没被这一轮连坐删掉）')
check(state()['existed'] == 0, '⑪b existed 计数回退到 0', str(state()['existed']))
with m.app.test_client() as c:
    jf2 = c.post('/api/manual/session/finish').get_json() or {}
check(jf2.get('ok') and jf2.get('saved') and jf2.get('rows') == 1,
      '⑪b 剩 1 条新客户正常落盘（灰字行不进 Excel）', str(jf2))

section('⑪c 已入库但 Excel 没写成功 → 那种行删了要连库一起删（唯一会动库的分支）')
# 「失败不清数据」的副作用：入库排在写表之前，所以 Excel 没写成功时行其实已经在库里了。
# 这种行若还只删内存 → 又变回「假删除」（下次采到同站被判「已在客户库」永远采不回来）。
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={'industry': '食品', 'country': '意大利'})
    c.post('/api/collect', json={'url': 'https://orphan-row.it', 'title': 'Orphan Row Srl',
                                 'html': html_with_mail('orphan-row.it')},
           headers={'X-DP-Ext': '1'})
    _real = m._save_enrich_excel
    m._save_enrich_excel = lambda *a, **k: False       # 模拟「结果 Excel 被 Excel 打开着」
    jf = c.post('/api/manual/session/finish').get_json() or {}
    m._save_enrich_excel = _real
check(jf.get('ok') is False, '⑪c 前置：落盘失败（本轮数据保留）', str(jf))
check(m.db_has_customer('https://orphan-row.it'), '⑪c 前置：入库在前 → 这条其实已经在库里了')
st = state()
check(len(st['session_rows']) == 1 and st['session_rows'][0]['kind'] == 'added',
      '⑪c 前置：明细还留着这条（等用户重试）', str(len(st['session_rows'])))
bk_before3 = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
_save_interval = m.db._DELETE_BACKUP_MIN_INTERVAL
m.db._DELETE_BACKUP_MIN_INTERVAL = 0          # 让这次备份必生成，好断言
try:
    with m.app.test_client() as c:
        jo = c.post('/api/manual/session/delete',
                    json={'rid': st['session_rows'][0]['rid']}).get_json() or {}
finally:
    m.db._DELETE_BACKUP_MIN_INTERVAL = _save_interval
check(jo.get('ok') is True and jo.get('deleted') == 1, '⑪c 这条已在库里 → 删它必须连库一起删', str(jo))
check(bool(jo.get('backup')), '⑪c 动库就要留删前备份', str(jo.get('backup')))
bk_after3 = sorted(os.listdir(m.db._backups_dir())) if os.path.isdir(m.db._backups_dir()) else []
check(any(f.startswith('客户库_删除前_') for f in bk_after3 if f not in bk_before3),
      '⑪c 备份文件名带「客户库_删除前_」', str(bk_after3))
check(not m.db_has_customer('https://orphan-row.it'), '⑪c 已从客户库真删掉（不是假删除）')
check(state()['session_active'] is True, '⑪c 删空后本轮仍进行中（没被误关）')
with m.app.test_client() as c:
    jz2 = c.post('/api/manual/session/finish').get_json() or {}
check(jz2.get('ok') and jz2.get('saved') is False, '⑪c 删空后结束 → 空轮不落盘', str(jz2))


# ============ ⑫ 去重键一致性：website_key 必须恒等于 normalize_domain(website) ============
# 为什么这条不变量值得单独守：db_has_customer()（判断「要不要跳过」）**只**按 website_key 查，
# 而 db_upsert_customer()（写入）是「按键查不到、再按 email 兜底合并」——两条口径本来就不一致。
# 键一旦与网址错位，那个网址就被永久判成「新客户」：每采一次往结果 Excel 多写一行，
# 而库里始终只有一条（实测：megazone.mx 在一份表里出现 3 次、另一份 2 次）。
section('⑫ 去重键自愈：website_key 恒等于 normalize_domain(website)')


def _conn():
    return m.db_connect()


def _one(sql, args=()):
    c = _conn()
    try:
        r = c.execute(sql, args).fetchone()
        return dict(r) if r is not None else None
    finally:
        c.close()


def _write(sql, args=()):
    c = _conn()
    try:
        cur = c.execute(sql, args)
        c.commit()
        return cur.lastrowid
    finally:
        c.close()


# 前置：一条记录，邮箱是两家站共用的（模拟「先采到 A 站，站上留的是 B 站邮箱」）
m.db_upsert_customer({'company_name': 'Old Key Ltd', 'website': 'https://keyshare-old.it',
                      'email': 'office@shared-key.it', 'source': '手动采集'})
check(m.db_has_customer('https://keyshare-old.it'), '⑫ 前置：A 站按老键入库')
check(m.db_has_customer('https://keyshare-new.it') is False, '⑫ 前置：B 站此刻还查不到')

# 再 upsert B 站（同邮箱）→ 走 email 兜底合并；键必须跟着网址一起改掉，不能留在旧站域名上
m.db_upsert_customer({'company_name': 'New Key Ltd', 'website': 'https://keyshare-new.it',
                      'email': 'office@shared-key.it', 'source': '手动采集'})
check(m.db_has_customer('https://keyshare-new.it'), '⑫ 自愈：合并后 B 站查得到了（键跟着网址走）')
check(m.db_has_customer('https://keyshare-old.it') is False, '⑫ 自愈：老键已让位')
r = _one("SELECT id, website_key FROM customers WHERE lower(email)='office@shared-key.it'")
check(r is not None and r['website_key'] == 'keyshare-new.it',
      '⑫ 库里仍只有一条，且键 = keyshare-new.it', str(r))
check(len(m.db_fix_website_keys(apply=False)[0]) == 0, '⑫ 自愈之后全库没有再错位')

# 端到端：自愈过的站再采一次 → 判「已在客户库」，不再算新增（这正是用户看到的现象）
with m.app.test_client() as c:
    c.post('/api/manual/session/start', json={'industry': '食品', 'country': '意大利'})
    jk = c.post('/api/collect', json={'url': 'https://keyshare-new.it', 'title': 'New Key Ltd',
                                      'html': html_with_mail('keyshare-new.it')},
                headers={'X-DP-Ext': '1'}).get_json() or {}
check(jk.get('existed') is True and jk.get('added') is False,
      '⑫ 端到端：自愈后再采 B 站 = 「已在客户库」，不再算新客户', str(jk))
stk = state()
check(stk['added'] == 0 and stk['existed'] == 1, '⑫ 本轮计数：新增 0 / 已在库 1', str(
    {'added': stk['added'], 'existed': stk['existed']}))
check(len(stk['session_rows']) == 1 and stk['session_rows'][0]['kind'] == 'existed',
      '⑫ 明细里只有一条 existed（它不会进 Excel）')

# 手工造一条错位（绕开 upsert）→ 扫描函数要能揪出来并修好
_write("UPDATE customers SET website_key='stale-bad-key.it' WHERE website_key='keyshare-new.it'")
fixed, conflict = m.db_fix_website_keys(apply=False)
check(len(fixed) == 1 and fixed[0][2] == 'keyshare-new.it' and not conflict,
      '⑫ dry-run 揪出 1 条错位、0 条冲突（且没写库）', str((fixed, conflict)))
check(_one("SELECT website_key FROM customers WHERE website_key='stale-bad-key.it'") is not None,
      '⑫ dry-run 确实只读：错位键还挂在那')
m.db_fix_website_keys(apply=True)
check(_one("SELECT id FROM customers WHERE website_key='keyshare-new.it'") is not None,
      '⑫ apply=True 真把键改回来了')
check(m.db_fix_website_keys(apply=False)[0] == [], '⑫ 修完再扫 → 不变量成立')

# 目标键已被**别的行**占用 → 只报冲突，绝不硬改（硬改会撞 UNIQUE 把整批写挂掉）
_write("INSERT INTO customers (company_name,website,website_key,email,phone,contact_url,country,"
       "industry,score,rating,snippet,search_formula,source,notes) "
       "VALUES ('Conflict Row','https://keyshare-new.it','other-key.it','','','','','',NULL,"
       "'','','','手动采集','')")
f3, c3 = m.db_fix_website_keys(apply=False)
check(len(c3) == 1 and c3[0][1] == 'other-key.it', '⑫ 目标键被占用 → 报冲突而不硬改', str(c3))
try:
    m.db_fix_website_keys(apply=True)
    check(True, '⑫ apply 模式下遇到冲突也不炸（不会撞 UNIQUE）')
except Exception as ex:
    check(False, '⑫ apply 模式下遇到冲突不炸', f'{type(ex).__name__}: {ex}')
check(_one("SELECT id FROM customers WHERE website_key='other-key.it'") is not None,
      '⑫ 冲突行保持原样，等人工判断')

print('\n' + '=' * 62)
print(f'通过 {_ok} 项，失败 {len(_fail)} 项')
for d in _fail:
    print('  ❌', d)
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if _fail else 0)
