# -*- coding: utf-8 -*-
"""扩展 ↔ 工具 的端到端联调（真 HTTP，不是 test_client）。

做法：用临时数据库/临时 output 起一个**真 Flask**（端口从 PREFERRED_PORTS 里挑一个能 bind 的，
因为扩展的端口发现只认这批候选），再把端口喂给 `_smoke_ext_bg.js` ——
那个脚本里跑的是**真实的 extension/background.js**，只把浏览器侧（chrome.* / fetch 出口）打了桩。

于是这一轮就等价于证明了「扩展 ↔ 工具」的接口契约成立：
  · /api/ping 的身份字段与端口发现优先级
  · /api/collect 的 X-DP-Ext 门禁、payload 形状、以及后端 collect_from_html 的返回结构
  · 采集结果真的落进了客户库（source='手动采集'）；node 侧的 ⑨ 用例再走一遍
    DP_MANUAL_STATE / DP_SESSION_DELETE / DP_SESSION_FINISH
    （即 popup 的「本轮采集明细」/ 行尾 ✕ 删除 / 「结束并生成 Excel」）：
    先删掉一条（明细 + 客户库双删），再把剩下的落成结果 Excel —— python 侧只复核结果，
    不再自己点一次
    ⚠️ 因为删掉了一条，所以这里的期望是「库里 1 条 / Excel 1 行」，不是 2。
  · 工具没在运行时的离线提示

外加三组只能用 test_client 验的（浏览器侧才讲 CORS 预检，够不着）：CORS 只回显
chrome-extension://、不带 X-DP-Ext 必 403、扩展心跳记账。

跑两遍完整用例：一遍让内容脚本缺席（走 executeScript 兜底），一遍让内容脚本在场
（验证「优先问内容脚本、不注入」）。两遍各用独立的临时目录，互不污染。

用法：python .workbuddy\\_smoke_ext_live.py
EXIT 0=全过  1=失败  3=跳过（本机找不到可用的候选端口）
"""
import importlib.util
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time

sys.stdout.reconfigure(encoding='utf-8')

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '客户开发工具.py')
BK = os.path.join(ROOT, '.workbuddy')
NODE = r'C:\Users\admin\.workbuddy\binaries\node\versions\22.22.2-3\node.exe'
HARNESS = os.path.join(BK, '_smoke_ext_bg.js')

_LOAD_N = [0]


def load_tool(tmpdir):
    """把主工具当成一个全新模块加载，并把库/结果文件都指向 tmpdir。"""
    _LOAD_N[0] += 1
    spec = importlib.util.spec_from_file_location(f'kf_ext{_LOAD_N[0]}', SRC)
    m = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = m
    spec.loader.exec_module(m)

    out = os.path.join(tmpdir, 'output')
    os.makedirs(out, exist_ok=True)
    m.db_path = lambda: os.path.join(tmpdir, 'test.sqlite3')
    m.db.db_path = m.db_path                     # ⚠️ 真正建连接的是 core/db.py 的 db_path()
    m.paths.ENRICH_FILE = os.path.join(out, 'result.xlsx')   # ⚠️ _append_result_row 读的是 core.paths
    m.ts_result_path = lambda industry='', country='': m.paths.ENRICH_FILE

    for p, what in ((m.db.db_path(), 'core.db.db_path'), (m.paths.ENRICH_FILE, 'ENRICH_FILE')):
        assert os.path.commonpath([os.path.realpath(p), os.path.realpath(tmpdir)]) == os.path.realpath(tmpdir), \
            f'{what} 没指向临时目录：{p}'

    # 明确禁止开浏览器：本轮的页面 HTML 由扩展侧提供，秒级路径就该命中
    def boom_enrich(page, url):
        raise AssertionError(f'不该开浏览器补全：{url}')

    m.enrich = boom_enrich
    m._manual_browser = boom_enrich
    m.create_browser_page = boom_enrich

    class NoNet:
        @staticmethod
        def get(*a, **k):
            raise OSError('stubbed network: 本轮不该走 requests 抓联系页')

    m.requests = NoNet()
    m.STATE['running'] = False
    m.MANUAL_STATE.update({'running': False, 'stop': False, 'ext_last_seen': 0.0})
    # ⚠️ core.db 是**同一个模块对象**（不像入口那样被 spec_from_file_location 重新加载），
    #    所以「删除前备份」的节流时间戳会从上一遍带过来 → 第二遍的删除会被节流成 backup=''，
    #    断言就会莫名其妙地红。这里显式清零，让每一遍都是干净的。
    m.db._last_delete_backup_at = 0.0
    m.init_customer_db()
    return m


def start_server(m, port):
    from werkzeug.serving import WSGIRequestHandler, make_server

    class Quiet(WSGIRequestHandler):
        """关掉 werkzeug 的每条访问日志：本轮的证据是 node 侧断言与工具自己的 log()，
        访问日志刷屏会把真正要看的东西冲走。"""
        def log_request(self, *a, **k):
            pass

        def log_message(self, *a, **k):
            pass

    m._RUN_PORT = port
    srv = make_server('127.0.0.1', port, m.app, threaded=True, request_handler=Quiet)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    return srv


def free_port():
    """拿一个「刚证明过没人用」的端口号（用于扮演失效的固定端口 / 冒名服务）。"""
    s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        s.bind(('127.0.0.1', 0))
        return s.getsockname()[1]
    finally:
        s.close()


def port_has_tool(port, timeout=0.7):
    """这个端口上是不是有本工具在跑（只认 /api/ping 的 app 身份字段）。"""
    import json
    import urllib.request
    try:
        with urllib.request.urlopen(f'http://127.0.0.1:{port}/api/ping', timeout=timeout) as r:
            d = json.loads(r.read().decode('utf-8'))
        return bool(d.get('ok') and d.get('app') == 'waimao-dp')
    except Exception:
        return False


def shadow_ports(port, cands):
    """候选表里排在测试端口**之前**、且真的跑着本工具的那些端口。

    ⚠️⚠️ 这不是可选的加固，是必须的：扩展按 PREFERRED_PORTS 顺序逐个试 /api/ping。
       本机 5000 落在 Windows 保留段（连不上，无害），但用户自己开着的那个实例常在 8000，
       而 8000 排在测试端口（如 8080）之前 —— 扩展会优先发现它。后果：
         ① 测试用的假网址会被真的写进**真实客户库**（实测踩过一次，多出 2 条假记录）；
         ② 所有端口断言都会失败。
       所以要把这些端口交给 harness，让它在 fetch 层当成连不上。
    """
    cands = list(cands)
    if port not in cands:
        return []
    return [p for p in cands[:cands.index(port)] if port_has_tool(p)]


def pick_preferred_port(cands):
    for p in cands:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            s.bind(('127.0.0.1', p))
            return p
        except OSError:
            continue
        finally:
            s.close()
    return None


def run_node(env_extra, timeout=180):
    env = dict(os.environ)
    env.update(env_extra)
    env['NODE_NO_WARNINGS'] = '1'
    r = subprocess.run([NODE, HARNESS], capture_output=True, text=True,
                       encoding='utf-8', errors='replace', env=env, timeout=timeout)
    print(r.stdout or '')
    if r.stderr and r.stderr.strip():
        print('--- node stderr ---')
        print(r.stderr.strip()[:2000])
    return r.returncode


FAILS = []


def check(cond, label, detail=''):
    if cond:
        print(f'  ✅ {label}')
    else:
        print(f'  ❌ {label}' + (f' —— {detail}' if detail else ''))
        FAILS.append(label)


def main():
    # 候选端口取自被加载的模块（_check_ext.py 已经保证它与 extension/ports.js 一致）
    spec = importlib.util.spec_from_file_location('kf_probe', SRC)
    probe = importlib.util.module_from_spec(spec)
    sys.modules['kf_probe'] = probe
    spec.loader.exec_module(probe)
    cands = list(probe.PREFERRED_PORTS)

    port = pick_preferred_port(cands)
    if not port:
        print(f'⚠️ 候选端口 {cands} 在本机一个都绑不上（Windows 保留段漂移），'
              f'扩展的端口发现没法真跑 → 跳过本轮联调。')
        print('   可临时改 PREFERRED_PORTS 后重跑；这不是代码缺陷。')
        return 3
    print(f'本轮用候选端口 {port} 起真 Flask（PREFERRED_PORTS={cands}）')

    # ⚠️ 环境里有别的实例在跑时必须遮蔽它，否则扩展会选中它，把测试假数据写进**真实客户库**
    shadows = shadow_ports(port, cands)
    if shadows:
        print(f'⚠️ 检测到本机 {shadows} 上跑着本工具的另一个实例（多半是你自己开着的那个）。')
        print(f'   它在候选表里排在测试端口 {port} 之前，而扩展是按顺序发现的 —— 会优先选中它，')
        print('   把本轮测试的假网址写进真实客户库。已把这些端口作为「影子端口」交给 harness，')
        print('   在 fetch 层当它不可达（等价于本次测试期间该实例不存在）。')
    else:
        print('本机候选端口上没有其它本工具实例在跑（无需遮蔽）')
    common_env = {'DP_TEST_PORT': str(port), 'DP_SHADOW_PORTS': ','.join(map(str, shadows))}

    stale = free_port()
    decoy = free_port()
    print(f'扮演「失效的固定端口」={stale}；扮演「冒名的别的服务」={decoy}')

    rc_all = 0

    # ---------------- 第 1 遍：内容脚本缺席（走 executeScript 兜底） ----------------
    print('\n' + '#' * 74)
    print('# 第 1 遍：内容脚本缺席 → background 用 executeScript 兜底读页面')
    print('#' * 74)
    tmp1 = tempfile.mkdtemp(prefix='ext_live1_')
    m1 = load_tool(tmp1)
    srv1 = start_server(m1, port)
    try:
        rc = run_node(dict(common_env, **{
            'DP_MOCK_PORTS': str(decoy),
            'DP_STALE_PORT': str(stale),
            'DP_CONTENT_SCRIPT': '0',
        }))
        rc_all |= rc
        # ⚠️ 手动采集现在是「本轮结束才落盘」：扩展只把行攒进本轮明细，不碰结果文件。
        #    「结束本轮」这一步**由 node 侧的 ⑨ 用例自己调 DP_SESSION_FINISH 完成**
        #    （等价于用户点面板上的「🏁 结束并生成 Excel」），所以这里只做**只读校验** ——
        #    再点一次会拿到 {ok:false, '当前没有进行中的本轮采集，无需结束'}。
        print('\n---------- 复核 node 侧那次「结束本轮」的结果（只读） ----------')
        with m1.app.test_client() as c:
            st1 = c.get('/api/manual/state').get_json() or {}
            print('   结束后的状态：', {k: st1.get(k) for k in
                                     ('session_active', 'session_rows', 'session_excel', 'session_error')})
            check(st1.get('session_active') is False, 'node 侧「结束本轮」后本轮应关闭')
            check(st1.get('session_rows') == [], '结束后本轮明细应清空')
            check(bool(st1.get('session_excel')), '结束后应记下生成的文件名',
                  str(st1.get('session_excel')))
            check(not st1.get('session_error'), '落盘不应留下错误标记', str(st1.get('session_error')))
    finally:
        srv1.shutdown()
    time.sleep(0.4)          # 等第 1 遍的 socket 彻底放开（两遍共用同一个端口）

    print('\n---------- 第 1 遍：库与结果文件确实是扩展那次 HTTP 调用写出来的 ----------')
    import sqlite3
    from openpyxl import load_workbook
    conn = sqlite3.connect(m1.db_path())
    conn.row_factory = sqlite3.Row
    rows = conn.execute('SELECT company_name,website_key,email,score,rating,industry,country,'
                        'source,search_formula FROM customers ORDER BY score DESC').fetchall()
    conn.close()
    for r in rows:
        print('   库里：', dict(r))
    check(len(rows) == 1, f'客户库里剩 1 条（node 侧 ⑨ 删掉了 1 条，实际 {len(rows)}）')
    check(all(r['source'] == '手动采集' for r in rows), 'source 全部是「手动采集」')
    check(all((r['search_formula'] or '').startswith('[手动采集]') for r in rows),
          'search_formula 全部带 [手动采集] 前缀（飞书会据此豁免低分门槛）')
    keys = sorted(r['website_key'] for r in rows)
    check(keys == ['plasticbagpc.com'], f'被删的那条真从库里消失，去重键正确：{keys}')
    check(any(r['industry'] == '塑料包装袋' for r in rows), '手选的行业写进了库')
    check(any(r['country'] == '德国' for r in rows), 'opts.country 写进了库')

    check(os.path.exists(m1.paths.ENRICH_FILE), '结果 Excel 被建出来了')
    if os.path.exists(m1.paths.ENRICH_FILE):
        wb = load_workbook(m1.paths.ENRICH_FILE)
        ws = wb.active
        head = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        data = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[1]]
        wb.close()
        check(list(head)[:10] == m1.EXCEL_HEADER_10, '结果 Excel 仍是 10 列硬契约', str(head))
        check(len(data) == 1, f'Excel 里 1 行数据（被删那条不进 Excel，实际 {len(data)}）')
        check(all(r[8].startswith('[手动采集]') for r in data), 'Excel 的「搜索公式」列带前缀')

    # ---------------- 第 2 遍：内容脚本在场 → 不注入 ----------------
    print('\n' + '#' * 74)
    print('# 第 2 遍：内容脚本在场 → 优先问内容脚本，不调 executeScript')
    print('#' * 74)
    tmp2 = tempfile.mkdtemp(prefix='ext_live2_')
    m2 = load_tool(tmp2)
    srv2 = start_server(m2, port)
    try:
        rc = run_node(dict(common_env, **{
            'DP_MOCK_PORTS': str(decoy),
            'DP_STALE_PORT': str(stale),
            'DP_CONTENT_SCRIPT': '1',
        }))
        rc_all |= rc
    finally:
        srv2.shutdown()

    # ---------------- 第 3 遍：工具没在运行 ----------------
    print('\n' + '#' * 74)
    print('# 第 3 遍：工具没在运行（端口全空）')
    print('#' * 74)
    tmp3 = tempfile.mkdtemp(prefix='ext_live3_')
    load_tool(tmp3)
    rc = run_node(dict(common_env, **{'DP_DEAD': '1', 'DP_DEAD_ONLY': '1'}))
    rc_all |= rc

    # ---------------- CORS / 门禁（只能 test_client 验，Node 侧没有 CORS） ----------------
    print('\n' + '#' * 74)
    print('# CORS 与门禁（浏览器才做预检，Node 够不着，这里用 test_client 补）')
    print('#' * 74)
    with m2.app.test_client() as c:
        pre = c.options('/api/collect', headers={
            'Origin': 'chrome-extension://abcdefghijklmnopabcdefghijklmnop',
            'Access-Control-Request-Method': 'POST',
            'Access-Control-Request-Headers': 'content-type,x-dp-ext',
        })
        check(pre.status_code == 204, '扩展来源的 OPTIONS 预检 → 204', str(pre.status_code))
        check(pre.headers.get('Access-Control-Allow-Origin', '').startswith('chrome-extension://'),
              '预检回显了 chrome-extension:// 来源',
              pre.headers.get('Access-Control-Allow-Origin', '<无>'))
        check('X-DP-Ext' in (pre.headers.get('Access-Control-Allow-Headers') or ''),
              'Allow-Headers 里含 X-DP-Ext',
              pre.headers.get('Access-Control-Allow-Headers', '<无>'))

        evil = c.options('/api/collect', headers={
            'Origin': 'https://evil.example.com',
            'Access-Control-Request-Method': 'POST',
        })
        check(not evil.headers.get('Access-Control-Allow-Origin'),
              '陌生来源的预检**没有** ACAO（浏览器会判失败，真请求发不出去）',
              evil.headers.get('Access-Control-Allow-Origin', '<无>'))
        check(evil.headers.get('Vary') == 'Origin', 'Vary: Origin（防中间缓存串源）')

        m2.MANUAL_STATE['ext_last_seen'] = 0.0
        hb = c.get('/api/ping', headers={'Origin': 'chrome-extension://abcdefghijklmnopabcdefghijklmnop'})
        check(hb.status_code == 200 and hb.get_json().get('app') == 'waimao-dp',
              '/api/ping 身份标记正确', str(hb.get_json()))
        check(m2.MANUAL_STATE['ext_last_seen'] > 0,
              '扩展的请求会刷新心跳（/manual 页据此显示「扩展在线」）',
              'ext_last_seen=' + str(m2.MANUAL_STATE['ext_last_seen']))

        bad = c.post('/api/collect', json={'url': 'https://x.com'})
        check(bad.status_code == 403, '不带 X-DP-Ext 的 POST → 403', str(bad.status_code))
        bad2 = c.post('/api/collect', data='url=https://x.com',
                      content_type='application/x-www-form-urlencoded',
                      headers={'X-DP-Ext': '1'})
        check(bad2.status_code == 200 and bad2.get_json().get('ok') is False,
              '非 JSON（简单请求）拿不到数据 → ok=false', str(bad2.get_json()))

    for t in (tmp1, tmp2, tmp3):
        shutil.rmtree(t, ignore_errors=True)

    print('\n' + '=' * 74)
    if FAILS or rc_all:
        print(f'❌ 失败：node rc={rc_all}，python 断言失败 {len(FAILS)} 项')
        for f in FAILS:
            print('   · ' + f)
        return 1
    print('✅ 扩展 ↔ 工具 端到端联调全部通过（真 HTTP，两遍路径 + 离线态 + CORS 门禁）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
