# -*- coding: utf-8 -*-
"""生成导入飞书的 JSON：邮箱只取第一个（email 字段只接受单个邮箱）"""
import json
import re
from openpyxl import load_workbook

wb = load_workbook('客户结果_含联系方式与评分.xlsx')
ws = wb.active
EMAIL_OK = re.compile(r'^[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}$')
records = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not (r and r[1] and r[5] and r[5] >= 50):
        continue
    title, url, emails, phones, contact, score, rating, snip, formula = r[:9]
    mails = [m for m in (emails or '').split(' / ') if EMAIL_OK.match(m.strip())]
    note = '脚本评分%d分（%s），待人工核实' % (score, rating)
    if len(mails) > 1:
        note += '；其他邮箱：' + ' / '.join(mails[1:])
    rec = {
        '公司名称': title,
        '官网': url,
        '电话': phones or '',
        '来源渠道': ['谷歌搜索'],
        '国家/地区': ['美国'],
        '使用的搜索公式': formula,
        '跟进状态': ['未联系'],
        '备注': note,
    }
    if mails:
        rec['邮箱'] = mails[0]
    records.append(rec)
with open('import_records.json', 'w', encoding='utf-8') as f:
    json.dump({'create_records': records}, f, ensure_ascii=False)
print('待导入记录数:', len(records))
